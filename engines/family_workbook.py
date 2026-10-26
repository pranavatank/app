"""Read-only household planning data from a ledger-style family workbook.

Source cells and cached formula values are retained. A summary FD balance
never becomes a made-up individual FD or a fictitious maturity receipt.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
import hashlib
import math
from pathlib import Path
import re

from engines.fd_planner_math import fy_start, interest_schedule


BANK_ALIASES = {
    'unity':'Unity Small Finance Bank', 'unitysmallfinancebank':'Unity Small Finance Bank',
    'jana':'Jana Small Finance Bank', 'janabank':'Jana Small Finance Bank',
    'janasmallfinancebank':'Jana Small Finance Bank',
    'ujjivan':'Ujjivan Small Finance Bank', 'ujjivansmallfinancebank':'Ujjivan Small Finance Bank',
    'equitas':'Equitas Small Finance Bank', 'equitassmallfinancebank':'Equitas Small Finance Bank',
    'idfc':'IDFC FIRST Bank', 'idfcfirst':'IDFC FIRST Bank', 'idfcfirstbank':'IDFC FIRST Bank',
    'hdfc':'HDFC Bank', 'hdfcbank':'HDFC Bank',
    'sbi':'State Bank of India', 'statebankofindia':'State Bank of India',
    'yes':'YES Bank', 'yesbank':'YES Bank', 'au':'AU Small Finance Bank',
    'ausmallfinancebank':'AU Small Finance Bank', 'boi':'Bank of India', 'bankofindia':'Bank of India',
    'rnsb':'Rajkot Nagarik Sahakari Bank', 'rajkotnagariksahakaribank':'Rajkot Nagarik Sahakari Bank',
}


def bank_name(text):
    key = re.sub(r'[^a-z0-9]', '', str(text or '').lower())
    key = re.sub(r'(limited|ltd)$', '', key)
    return BANK_ALIASES.get(key, str(text or '').strip())


def number(value):
    if value is None or value == '':
        return None
    if isinstance(value, str):
        value = re.sub(r'[₹,\s]', '', value)
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    return result if math.isfinite(result) and result >= 0 else None


def day(value):
    if isinstance(value, datetime): return value.date()
    if isinstance(value, date): return value
    if isinstance(value, str):
        try: return date.fromisoformat(value[:10])
        except ValueError: pass
    return None


def stable_id(text):
    return int(hashlib.sha256(text.encode()).hexdigest()[:12], 16)


@dataclass
class FamilyData:
    people: list[dict] = field(default_factory=list)
    accounts: list[dict] = field(default_factory=list)
    deposits: list[dict] = field(default_factory=list)
    receipts: list[dict] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)
    source_path: str = ''
    source_hash: str = ''
    source_modified: str = ''
    financial_year: int = 2026
    tax_complete: bool = True

    @property
    def cash(self):
        return sum(a['cash'] for a in self.accounts if a['eligible'])

    @property
    def banks(self):
        return sorted({a['bank'] for a in self.accounts})


def read_family_workbook(path, as_of=None, ledger_people=(), ledger_accounts=(), ledger_fds=(), ledger_expected=()):
    import openpyxl
    as_of = as_of or date.today()
    path = Path(path)
    data = FamilyData(source_path=str(path.resolve()), source_hash=hashlib.sha256(path.read_bytes()).hexdigest(),
                      source_modified=datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec='seconds'),
                      financial_year=fy_start(as_of))
    values = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        if 'Bank' not in values.sheetnames:
            raise ValueError('Expected a Bank worksheet with bank/holder account rows.')
        sheet = values['Bank']
        records = list(sheet.iter_rows(values_only=True))
        def cell(row, col):
            return records[row-1][col-1] if 0 < row <= len(records) and col <= len(records[row-1]) else None
        header = next((i for i in range(1,len(records)+1)
                       if str(cell(i,6) or '').strip().lower() == 'saving balance'
                       and 'fd' in str(cell(i,7) or '').lower()), None)
        if header is None:
            raise ValueError('Could not find the account summary header (Saving Balance / No. of FD).')
        ledger_names = {str(p.get('full_name','')).strip().casefold():p for p in ledger_people}
        people = {}
        current_bank = None
        for row in range(header+1, len(records)+1):
            raw_bank, holder = cell(row,3), cell(row,4)
            if raw_bank and str(raw_bank).strip().lower() in BANK_ALIASES:
                current_bank = bank_name(raw_bank)
            elif raw_bank and not holder:
                current_bank = None
            if current_bank is None or not isinstance(holder,str) or not holder.strip():
                continue
            holder=holder.strip()
            joint='+' in holder or '&' in holder
            pid=stable_id('person:'+holder.casefold())
            if pid not in people:
                ledger=ledger_names.get(holder.casefold(),{})
                born=day(ledger.get('date_of_birth'))
                if born and (born>as_of or (as_of-born).days<18*365):
                    data.notes.append(f'{holder}: birth date needs review; general rates only.')
                    born=None
                kind='huf' if holder.casefold()=='huf' else ('joint' if joint else 'individual')
                people[pid]=dict(id=pid,name=holder,birth_date=born,tax_kind=kind,income_by_fy={},
                                 tax_complete=True,ledger_id=ledger.get('person_id'))
            balance=number(cell(row,6))
            matching=[a for a in ledger_accounts if bank_name(a.get('bank_name'))==current_bank
                      and a.get('person_id')==people[pid].get('ledger_id')
                      and (a.get('account_type') or '').casefold() in ('savings','salary','current')]
            balance_source=f'Bank!F{row}'
            if balance is None and len(matching)==1:
                balance=number(matching[0].get('current_balance')); balance_source='App account balance'
            count=number(cell(row,7)); principal=number(cell(row,8)); maturity=number(cell(row,9))
            account=dict(id=stable_id(f'account:{holder.casefold()}:{current_bank}'),person_id=pid,
                         person=holder,bank=current_bank,label=f'{current_bank} · {holder}',
                         cash=balance or 0, cash_known=balance is not None, cash_source=balance_source,
                         fd_count=int(count) if count is not None else None,fd_principal=principal,
                         fd_maturity=maturity,existing_fd=0.0,unknown_fds=0,
                         eligible=not joint,source_row=row,source_sheet='Bank',details=[])
            if principal is not None and count is None:
                account['fd_count']=1
            data.accounts.append(account)
        if not data.accounts:
            raise ValueError('No recognised bank/holder account rows were found.')
        by_holder_bank={(a['person'].casefold(),a['bank']):a for a in data.accounts}
        current_account=None
        for row in range(1,len(records)+1):
            title=cell(row,3)
            if isinstance(title,str) and '-' in title:
                left,right=title.split('-',1)
                current_account=by_holder_bank.get((right.strip().casefold(),bank_name(left)))
            start=day(cell(row,3)); days=number(cell(row,4)); principal=number(cell(row,6)); rate=number(cell(row,7))
            if not current_account or not start or not days or not principal or not rate:
                continue
            end=day(cell(row,5)) or start+timedelta(days=int(days))
            if end<=start or not 0<rate<=20:
                data.notes.append(f'Bank row {row}: invalid FD dates or coupon; excluded from projected receipts.')
                continue
            maturity=number(cell(row,10))
            schedule=interest_schedule(start,(end-start).days,rate,4,'actual')
            estimated=principal*(1+sum(schedule.values()))
            deposit=dict(id=f'Bank:{row}',person_id=current_account['person_id'],person=current_account['person'],
                         account_id=current_account['id'],bank=current_account['bank'],principal=principal,
                         rate=rate,start=start.isoformat(),maturity_date=end.isoformat(),
                         maturity_amount=maturity if maturity is not None else estimated,
                         maturity_is_estimate=True,active=end>as_of,interest_by_fy=schedule,
                         source_cell=f'Bank!C{row}:J{row}')
            data.deposits.append(deposit); current_account['details'].append(deposit)
            if end>as_of:
                data.receipts.append(dict(date=end.isoformat(),person=deposit['person'],bank=deposit['bank'],
                                          amount=deposit['maturity_amount'],kind='FD maturity',confirmed=False))
            else:
                data.notes.append(f'{deposit["person"]} · {deposit["bank"]}: a listed FD has already matured. Its proceeds are not added to current cash.')
        for account in data.accounts:
            # Supplement details with complete ledger FDs, deduplicated by
            # bank/owner/start/principal/maturity rather than account numbers.
            known={(d['start'],d['principal'],d['maturity_date']) for d in account['details']}
            ledger_pid=people[account['person_id']].get('ledger_id')
            for fd in ledger_fds:
                if (fd.get('person_id')!=ledger_pid or bank_name(fd.get('bank_name'))!=account['bank']
                        or fd.get('status')!='Active'):
                    continue
                start,end=day(fd.get('start_date')),day(fd.get('maturity_date'))
                principal,rate=number(fd.get('principal_amount')),number(fd.get('interest_rate'))
                if not start or not end or end<=max(start,as_of) or not principal or not rate:continue
                key=(start.isoformat(),principal,end.isoformat())
                if key in known:continue
                schedule=interest_schedule(start,(end-start).days,rate)
                deposit=dict(id=f'ledger:{fd["fd_id"]}',person_id=account['person_id'],person=account['person'],
                             account_id=account['id'],bank=account['bank'],principal=principal,rate=rate,
                             start=start.isoformat(),maturity_date=end.isoformat(),
                             maturity_amount=max(principal,number(fd.get('maturity_amount')) or principal*(1+sum(schedule.values()))),
                             maturity_is_estimate=True,active=True,interest_by_fy=schedule,source_cell='App FD record')
                account['details'].append(deposit);data.deposits.append(deposit);known.add(key)
                data.receipts.append(dict(date=end.isoformat(),person=account['person'],bank=account['bank'],
                                          amount=deposit['maturity_amount'],kind='FD maturity',confirmed=False))
            active=[d for d in account['details'] if d['active']]
            detail_principal=sum(d['principal'] for d in active)
            summary_principal=account['fd_principal'] or 0
            account['unknown_fds']=max(0,(account['fd_count'] or 0)-len(active))
            unmatched=max(0,summary_principal-detail_principal)
            account['existing_fd']=max(account['fd_maturity'] or 0,
                                       sum(d['maturity_amount'] for d in active)+unmatched)
            account['terms_complete']=account['unknown_fds']==0 and unmatched<.01
            if not account['terms_complete']:
                data.tax_complete=False;people[account['person_id']]['tax_complete']=False
            if not account['cash_known']:
                data.notes.append(f'{account["person"]} · {account["bank"]}: savings balance is missing; not counted as investable cash.')
            for d in active:
                for year,coefficient in d['interest_by_fy'].items():
                    member=people[d['person_id']]
                    member['income_by_fy'][year]=member['income_by_fy'].get(year,0)+d['principal']*coefficient
        salary_months=set()
        if 'Salary' in values.sheetnames:
            salary=list(values['Salary'].iter_rows(values_only=True))
            current_person=None
            for row in salary:
                label=row[2] if len(row)>2 else None
                match=next((p for p in people.values() if isinstance(label,str) and p['name'].casefold()==label.strip().casefold()),None)
                if match: current_person=match
                payment_day=day(row[4]) if len(row)>4 else None
                expected=number(row[5]) if len(row)>5 else None
                actual=number(row[6]) if len(row)>6 else None
                if current_person and payment_day and (expected is not None or actual is not None):
                    salary_months.add((current_person['id'],payment_day.strftime('%Y-%m')))
                    year=fy_start(payment_day)
                    current_person['income_by_fy'][year]=current_person['income_by_fy'].get(year,0)+(actual if actual is not None else expected)
                    if payment_day>as_of and expected is not None and expected>0 and actual is None:
                        data.receipts.append(dict(date=payment_day.isoformat(),person=current_person['name'],bank='',
                                                  amount=expected,kind='Expected salary',confirmed=False))
        # App forecasts supplement the source, without adding a second salary
        # forecast for a workbook month or reusing a matched actual receipt.
        by_ledger={p['ledger_id']:p for p in people.values() if p.get('ledger_id') is not None}
        for row in ledger_expected:
            member=by_ledger.get(row.get('person_id'))
            payment_day=day(row.get('expected_date'));amount=number(row.get('expected_amount'))
            category=str(row.get('income_type') or 'income').strip()
            kind=category.casefold()
            if (not member or not payment_day or amount is None or row.get('actual_transaction_id') is not None
                    or kind in ('fd interest','savings interest')):continue
            if 'salary' in kind and (member['id'],payment_day.strftime('%Y-%m')) in salary_months:continue
            year=fy_start(payment_day)
            member['income_by_fy'][year]=member['income_by_fy'].get(year,0)+amount
            if payment_day>as_of and amount>0:
                data.receipts.append(dict(date=payment_day.isoformat(),person=member['name'],bank=row.get('bank_name') or '',
                                          amount=amount,kind=f'Expected {category} (app)',confirmed=False))
        data.people=list(people.values())
        if not data.tax_complete:
            data.notes.insert(0,'Some existing FDs have aggregate balances only. Their unknown interest is not treated as zero; tax/headroom estimates are provisional.')
        data.notes.append('Savings balances are a workbook/app snapshot. Future salary and FD maturities are shown separately and are not cash available today.')
        data.notes.append('App income expectations supplement workbook forecasts; matched receipts and overlapping salary months are not counted twice. Missing future income still needs review.')
        data.notes.append('Family transfers are planning scenarios. Tax remains subject to beneficial ownership, spouse/minor clubbing and HUF rules.')
        data.receipts.sort(key=lambda row:row['date'])
        return data
    finally:
        values.close()
