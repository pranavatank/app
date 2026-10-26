from datetime import date,datetime
import hashlib
import pytest
from engines.family_workbook import read_family_workbook,bank_name


@pytest.fixture
def workbook(tmp_path):
    from openpyxl import Workbook
    book=Workbook();sheet=book.active;sheet.title='Bank'
    sheet.cell(5,6,'Saving Balance');sheet.cell(5,7,'No. of FD')
    for r,bank,person,cash,count,principal in [(6,'Unity','A',100000,0,None),(7,None,'B',0,0,None),
                                              (9,'Jana','A',1000,2,20000),(10,None,'HUF',500,0,None),
                                              (12,'HDFC','B',None,0,None)]:
        for col,value in [(3,bank),(4,person),(6,cash),(7,count),(8,principal)]:
            if value is not None:sheet.cell(r,col,value)
    sheet.cell(16,3,'JANA-A');sheet.cell(17,3,datetime(2026,1,1));sheet.cell(17,4,365)
    sheet.cell(17,5,datetime(2027,1,1));sheet.cell(17,6,10000);sheet.cell(17,7,8);sheet.cell(17,10,10800)
    salary=book.create_sheet('Salary');salary.cell(1,3,'A')
    salary.cell(2,5,datetime(2026,11,1));salary.cell(2,6,10000)
    path=tmp_path/'family.xlsx';book.save(path);return path


def test_workbook_source_is_unchanged_and_cash_is_not_future_income(workbook):
    before=hashlib.sha256(workbook.read_bytes()).hexdigest()
    data=read_family_workbook(workbook,date(2026,10,26))
    assert data.cash==101500
    assert len(data.accounts)==5
    assert len(data.receipts)==2
    assert hashlib.sha256(workbook.read_bytes()).hexdigest()==before


def test_aggregate_fd_summary_does_not_create_fake_maturities(workbook):
    data=read_family_workbook(workbook,date(2026,10,26))
    jana=next(a for a in data.accounts if a['bank']=='Jana Small Finance Bank' and a['person']=='A')
    assert jana['unknown_fds']==1
    assert jana['existing_fd']==20800
    assert len(data.deposits)==1
    assert not data.tax_complete


def test_missing_cash_and_zero_cash_have_different_statuses(workbook):
    data=read_family_workbook(workbook,date(2026,10,26))
    hdfc=next(a for a in data.accounts if a['bank']=='HDFC Bank')
    zero=next(a for a in data.accounts if a['bank']=='Unity Small Finance Bank' and a['person']=='B')
    assert hdfc['cash']==zero['cash']==0
    assert not hdfc['cash_known']
    assert zero['cash_known']


def test_huf_is_not_classified_as_an_individual(workbook):
    data=read_family_workbook(workbook,date(2026,10,26))
    assert next(p for p in data.people if p['name']=='HUF')['tax_kind']=='huf'


def test_correct_ledger_dates_are_used_and_bad_dates_do_not_gain_senior_rates(workbook):
    people=[dict(person_id=7,full_name='A',date_of_birth='1974-12-07'),
            dict(person_id=8,full_name='B',date_of_birth='2074-01-01')]
    data=read_family_workbook(workbook,date(2026,10,26),people)
    assert next(p for p in data.people if p['name']=='A')['birth_date']==date(1974,12,7)
    assert next(p for p in data.people if p['name']=='B')['birth_date'] is None


def test_app_forecasts_supplement_workbook_without_duplicate_salary(workbook):
    people=[dict(person_id=7,full_name='A',date_of_birth='1990-01-01')]
    def expected(day,amount,kind='Salary',matched=None):
        return dict(person_id=7,expected_date=day,expected_amount=amount,income_type=kind,actual_transaction_id=matched)
    rows=[expected('2026-11-05',10000),expected('2027-04-01',15000),expected('2026-12-05',3000,'Rent'),
          expected('2026-12-06',9000,'Salary',123),expected('2026-12-07',8000,'FD Interest')]
    data=read_family_workbook(workbook,date(2026,10,26),ledger_people=people,ledger_expected=rows)
    member=next(p for p in data.people if p['name']=='A')
    assert member['income_by_fy'][2027]==15000
    assert len([r for r in data.receipts if r['kind']=='Expected salary'])==1
    assert len([r for r in data.receipts if r['kind'].endswith('(app)')])==2
    assert data.cash==101500


@pytest.mark.parametrize('name,expected',[('Jana Bank','Jana Small Finance Bank'),('Unity','Unity Small Finance Bank'),
                                          ('EQUITAS SMALL FINANCE BANK LIMITED','Equitas Small Finance Bank'),('RNSB','Rajkot Nagarik Sahakari Bank')])
def test_bank_aliases(name,expected):assert bank_name(name)==expected


def test_actual_salary_is_counted_when_the_expected_cell_is_blank(workbook):
    from openpyxl import load_workbook
    book=load_workbook(workbook);book['Salary'].cell(2,6).value=None;book['Salary'].cell(2,7,11000)
    book.save(workbook);book.close()
    data=read_family_workbook(workbook,date(2026,10,26))
    person=next(p for p in data.people if p['name']=='A')
    existing=sum(d['principal']*d['interest_by_fy'].get(2026,0) for d in data.deposits if d['person']=='A')
    assert person['income_by_fy'][2026]==pytest.approx(11000+existing)
    assert not any(r['kind']=='Expected salary' for r in data.receipts)
